from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand

from openpyxl import load_workbook

from cartoes.models import Cartao, Fatura, ItemFatura
from despesas_basicas.models import DespesaBasica
from filhos.models import DespesaFilho, Filho
from imoveis_transporte.models import Consorcio, ParcelaConsorcio
from inicio.constants import ANO_PADRAO, STATUS_ABERTO, STATUS_CARTAO, STATUS_FECHADO
from inicio.models import Receita

MESES_NOME = {
    'janeiro': 1,
    'fevereiro': 2,
    'março': 3,
    'marco': 3,
    'abril': 4,
    'maio': 5,
    'junho': 6,
    'julho': 7,
    'agosto': 8,
    'setembro': 9,
    'outubro': 10,
    'novembro': 11,
    'dezembro': 12,
}

FONTES = {
    'salário aline': Receita.FONTE_SALARIO_ALINE,
    'salario aline': Receita.FONTE_SALARIO_ALINE,
    'salário luiz': Receita.FONTE_SALARIO_LUIZ,
    'salario luiz': Receita.FONTE_SALARIO_LUIZ,
    'mesada vovô': Receita.FONTE_MESADA_VOVO,
    'mesada vovo': Receita.FONTE_MESADA_VOVO,
    'outros': Receita.FONTE_OUTROS,
}


def _status(valor):
    if valor is None:
        return STATUS_ABERTO
    texto = str(valor).strip().lower()
    if texto == 'fechado':
        return STATUS_FECHADO
    if texto in {'cartão', 'cartao'}:
        return STATUS_CARTAO
    return STATUS_ABERTO


def _decimal(valor):
    if valor is None or valor == '':
        return None
    if isinstance(valor, str) and valor.startswith('='):
        return None
    try:
        return Decimal(str(valor))
    except (InvalidOperation, TypeError, ValueError):
        return None


def _texto(valor):
    if valor is None:
        return ''
    return str(valor).strip()


def _mes_por_nome(valor):
    return MESES_NOME.get(_texto(valor).lower())


class Command(BaseCommand):
    help = 'Importa os dados da planilha Despesas fixas e variáveis 2026.xlsx'

    def add_arguments(self, parser):
        parser.add_argument(
            '--arquivo',
            default=str(settings.PLANILHA_ORIGEM),
            help='Caminho da planilha .xlsx',
        )
        parser.add_argument('--ano', type=int, default=ANO_PADRAO)

    def handle(self, *args, **options):
        caminho = Path(options['arquivo'])
        if not caminho.exists():
            self.stderr.write(self.style.ERROR(f'Arquivo não encontrado: {caminho}'))
            return

        ano = options['ano']
        wb = load_workbook(caminho, data_only=True)
        self._importar_despesas_basicas(wb, ano)
        self._importar_filhos(wb, ano)
        self._importar_cartoes(wb, ano)
        self._importar_imoveis(wb, ano)
        self._importar_receitas(wb, ano)
        self.stdout.write(self.style.SUCCESS(f'Planilha importada para o ano {ano}.'))

    def _importar_despesas_basicas(self, wb, ano):
        ws = wb['Desp.básicas']
        blocos = [
            (1, list(range(2, 11))),
            (13, list(range(14, 23))),
            (25, list(range(26, 35))),
        ]
        colunas = [(1, 2, 3), (4, 5, 6), (7, 8, 9), (10, 11, 12)]
        criados = 0
        for linha_mes, linhas_dados in blocos:
            for col_desc, col_valor, col_status in colunas:
                mes = _mes_por_nome(ws.cell(linha_mes, col_desc).value)
                if not mes:
                    continue
                for linha in linhas_dados:
                    descricao = _texto(ws.cell(linha, col_desc).value)
                    if not descricao or descricao.lower() == 'total':
                        continue
                    valor = _decimal(ws.cell(linha, col_valor).value) or Decimal('0')
                    DespesaBasica.objects.update_or_create(
                        ano=ano,
                        mes=mes,
                        descricao=descricao,
                        defaults={
                            'valor': valor,
                            'status': _status(ws.cell(linha, col_status).value),
                        },
                    )
                    criados += 1
        self.stdout.write(f'Despesas básicas: {criados}')

    def _importar_filhos(self, wb, ano):
        ws = wb['Filhos']
        joao, _ = Filho.objects.get_or_create(nome='João Lucas')
        beatriz, _ = Filho.objects.get_or_create(nome='Beatriz Helena')
        blocos = [
            (2, list(range(3, 10))),
            (12, list(range(13, 20))),
            (22, list(range(23, 30))),
        ]
        grupos = [
            (joao, [(1, 2, 3), (4, 5, 6), (7, 8, 9), (10, 11, 12)]),
            (beatriz, [(13, 14, 15), (16, 17, 18), (19, 20, 21), (22, 23, 24)]),
        ]
        criados = 0
        for filho, colunas in grupos:
            for linha_mes, linhas_dados in blocos:
                for col_desc, col_valor, col_status in colunas:
                    mes = _mes_por_nome(ws.cell(linha_mes, col_desc).value)
                    if not mes:
                        continue
                    for linha in linhas_dados:
                        descricao = _texto(ws.cell(linha, col_desc).value)
                        if not descricao or descricao.lower() == 'total':
                            continue
                        valor = _decimal(ws.cell(linha, col_valor).value) or Decimal('0')
                        DespesaFilho.objects.update_or_create(
                            filho=filho,
                            ano=ano,
                            mes=mes,
                            descricao=descricao,
                            defaults={
                                'valor': valor,
                                'status': _status(ws.cell(linha, col_status).value),
                            },
                        )
                        criados += 1
        self.stdout.write(f'Despesas dos filhos: {criados}')

    def _importar_cartoes(self, wb, ano):
        ws = wb['Cartões']
        blocos = [
            {'linha_nome': 1, 'linha_mes': 2, 'inicio': 3, 'fim': 45, 'total': 46},
            {'linha_nome': 47, 'linha_mes': 48, 'inicio': 49, 'fim': 93, 'total': 94},
            {'linha_nome': 95, 'linha_mes': 96, 'inicio': 97, 'fim': 139, 'total': 140},
        ]
        lados = [
            {'nome_col': 1, 'meses': [(1, 2), (3, 4), (5, 6), (7, 8)], 'titular': 'Aline'},
            {'nome_col': 10, 'meses': [(10, 11), (12, 13), (14, 15), (16, 17)], 'titular': 'Luiz'},
        ]
        criados = 0
        for bloco in blocos:
            for lado in lados:
                nome_cartao = _texto(ws.cell(bloco['linha_nome'], lado['nome_col']).value)
                if not nome_cartao or nome_cartao.lower() in {'x', 'total'}:
                    nome_cartao = f"Nubank {lado['titular']}"
                cartao, _ = Cartao.objects.get_or_create(
                    nome=nome_cartao,
                    defaults={'titular': lado['titular']},
                )
                for col_desc, col_valor in lado['meses']:
                    mes = _mes_por_nome(ws.cell(bloco['linha_mes'], col_desc).value)
                    if not mes:
                        continue
                    fatura, _ = Fatura.objects.update_or_create(
                        cartao=cartao,
                        ano=ano,
                        mes=mes,
                        defaults={'status': _status(ws.cell(bloco['linha_mes'], col_valor).value)},
                    )
                    fatura.itens.all().delete()
                    for linha in range(bloco['inicio'], bloco['fim'] + 1):
                        descricao = _texto(ws.cell(linha, col_desc).value)
                        valor = _decimal(ws.cell(linha, col_valor).value)
                        if not descricao or valor is None:
                            continue
                        if descricao.lower() == 'total':
                            continue
                        ItemFatura.objects.create(
                            fatura=fatura,
                            descricao=descricao,
                            valor=valor,
                        )
                        criados += 1
                    if not fatura.itens.exists():
                        fatura.delete()
        self.stdout.write(f'Itens de cartão: {criados}')

    def _importar_imoveis(self, wb, ano):
        ws = wb['ImóveisTransporte']
        grupos = [
            (2, 3, 4),
            (5, 6, 7),
            (8, 9, 10),
            (11, 12, 13),
            (14, 15, 16),
            (17, 18, 19),
        ]
        criados = 0
        for col_parcela, col_valor, col_status in grupos:
            nome_completo = _texto(ws.cell(1, col_parcela).value)
            if not nome_completo or nome_completo.lower() == 'livre':
                continue
            nome, grupo, total_parcelas = self._parse_consorcio(nome_completo)
            consorcio, _ = Consorcio.objects.update_or_create(
                nome=nome_completo,
                defaults={'grupo': grupo, 'total_parcelas': total_parcelas},
            )
            for linha in range(3, 15):
                data = ws.cell(linha, 1).value
                if data is None:
                    continue
                mes = data.month if hasattr(data, 'month') else linha - 2
                rotulo = _texto(ws.cell(linha, col_parcela).value)
                valor = _decimal(ws.cell(linha, col_valor).value)
                if valor is None:
                    continue
                numero = None
                if '/' in rotulo:
                    try:
                        numero = int(rotulo.split('/')[0])
                        total_parcelas = total_parcelas or int(rotulo.split('/')[1])
                    except ValueError:
                        numero = None
                ParcelaConsorcio.objects.update_or_create(
                    consorcio=consorcio,
                    ano=ano,
                    mes=mes,
                    defaults={
                        'rotulo_parcela': rotulo,
                        'numero_parcela': numero,
                        'valor': valor,
                        'status': _status(ws.cell(linha, col_status).value),
                    },
                )
                criados += 1
            if total_parcelas:
                consorcio.total_parcelas = total_parcelas
                consorcio.save(update_fields=['total_parcelas'])
        self.stdout.write(f'Parcelas de consórcio: {criados}')

    def _parse_consorcio(self, nome_completo):
        grupo = ''
        if ' - ' in nome_completo:
            nome, resto = nome_completo.split(' - ', 1)
            grupo = resto.replace('Grupo ', '').strip()
        else:
            nome = nome_completo
        total = None
        return nome.strip(), grupo, total

    def _importar_receitas(self, wb, ano):
        ws = wb['Início 2025']
        blocos = [
            {'meses': [(1, 2), (4, 5), (7, 8), (10, 11)], 'linhas': range(11, 15)},
            {'meses': [(1, 2), (4, 5), (7, 8), (10, 11)], 'linhas': range(33, 37)},
            {'meses': [(1, 2), (4, 5), (7, 8), (10, 11)], 'linhas': range(55, 59)},
        ]
        meses_bloco = [
            [1, 2, 3, 4],
            [5, 6, 7, 8],
            [9, 10, 11, 12],
        ]
        criados = 0
        for bloco, meses in zip(blocos, meses_bloco):
            for (col_nome, col_valor), mes in zip(bloco['meses'], meses):
                for linha in bloco['linhas']:
                    fonte_nome = _texto(ws.cell(linha, col_nome).value)
                    fonte = FONTES.get(fonte_nome.lower())
                    valor = _decimal(ws.cell(linha, col_valor).value)
                    if not fonte or valor is None:
                        continue
                    Receita.objects.update_or_create(
                        ano=ano,
                        mes=mes,
                        fonte=fonte,
                        defaults={'valor': valor},
                    )
                    criados += 1
        self.stdout.write(f'Receitas: {criados}')
