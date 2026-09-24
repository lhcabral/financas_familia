import csv
import io
import re
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook

from cartoes.models import Fatura, ItemFatura
from inicio.constants import STATUS_ABERTO
from inicio.recorrencia import avancar_mes

COLUNAS_DESCRICAO = {
    'descricao', 'descrição', 'desc', 'item', 'estabelecimento',
    'historico', 'histórico', 'histуrico', 'title', 'titulo', 'título',
}
COLUNAS_VALOR = {
    'valor', 'value', 'amount', 'preco', 'preço',
    'valor(r$)', 'valor (r$)', 'valor_rs', 'valor r$', 'débito', 'debito',
}
COLUNAS_QTD = {
    'quantidade_parcelas', 'qtd_parcelas', 'qtd parcelas', 'parcelas',
    'total_parcelas', 'n_parcelas', 'nº parcelas',
}
COLUNAS_ATUAL = {
    'parcela_atual', 'parcela atual', 'parcela', 'atual', 'n_parcela',
}

IGNORAR_DESCRICAO = re.compile(
    r'^(saldo anterior|pagto\.?\s*por\s*deb|pagamento(\s+recebido)?|total|resumo|'
    r'situa[cç][aã]o|data$|l[a-z].*;;;|cota[cç][aã]o)',
    re.IGNORECASE,
)
PARCELA_NO_TEXTO = re.compile(
    r'^(?P<desc>.+?)\s+(?P<atual>\d{1,2})\s*/\s*(?P<total>\d{1,2})\s*$',
)
PARCELA_NUBANK = re.compile(
    r'^(?P<desc>.+?)\s*[-–]\s*Parcela\s+(?P<atual>\d+)\s*/\s*(?P<total>\d+)\s*$',
    re.IGNORECASE,
)
PARCELA_NO_MEIO = re.compile(r'\b(?P<atual>\d{1,2})\s*/\s*(?P<total>\d{1,2})\b')
CIDADES_COMPOSTAS_INICIO = {
    'SAO', 'SÃO', 'RIO', 'BELO', 'PORTO', 'CABO', 'MONTE', 'CAMPO', 'SANTA', 'SANTO',
}


def salvar_item_com_parcelas(item):
    """Salva o item e lança as parcelas restantes nas faturas dos próximos meses."""
    with transaction.atomic():
        item.save()
        return propagar_parcelas(item)


def propagar_parcelas(item):
    """Cria itens nas faturas seguintes enquanto houver parcelas em aberto."""
    restantes = item.parcelas_restantes
    if restantes <= 0:
        return 0

    fatura = item.fatura
    criados = 0
    for passo in range(1, restantes + 1):
        ano, mes = avancar_mes(fatura.ano, fatura.mes, passo)
        proxima, _ = Fatura.objects.get_or_create(
            cartao=fatura.cartao,
            ano=ano,
            mes=mes,
            defaults={'status': STATUS_ABERTO},
        )
        parcela = (item.parcela_atual or 1) + passo
        if _item_parcela_existe(proxima, item.descricao, parcela, item.quantidade_parcelas):
            continue
        ItemFatura.objects.create(
            fatura=proxima,
            descricao=item.descricao,
            valor=item.valor,
            quantidade_parcelas=item.quantidade_parcelas,
            parcela_atual=parcela,
        )
        criados += 1
    return criados


def _item_parcela_existe(fatura, descricao, parcela_atual, quantidade_parcelas):
    return fatura.itens.filter(
        descricao=descricao,
        parcela_atual=parcela_atual,
        quantidade_parcelas=quantidade_parcelas,
    ).exists()


def importar_itens_arquivo(fatura, arquivo):
    """
    Importa itens para a fatura a partir de CSV/XLS/XLSX.
    Formatos: CSV simples, Bradesco, Nubank e fatura Porto Seguro (XLSX).
    """
    nome = (getattr(arquivo, 'name', '') or '').lower()
    if nome.endswith(('.xlsx', '.xlsm', '.xls')):
        linhas = list(_iter_linhas_planilha(arquivo))
    else:
        texto = _decodificar(arquivo)
        linhas = list(_iter_linhas_csv(texto))

    criados = 0
    propagados = 0
    erros = []

    for indice, linha in linhas:
        try:
            dados = _linha_para_item(linha)
        except ValueError as exc:
            erros.append(f'Linha {indice}: {exc}')
            continue
        if dados is None:
            continue
        item = ItemFatura(fatura=fatura, **dados)
        propagados += salvar_item_com_parcelas(item)
        criados += 1

    return criados, propagados, erros


# Compatibilidade com código/testes anteriores.
importar_itens_csv = importar_itens_arquivo


def _decodificar(arquivo):
    bruto = arquivo.read()
    if isinstance(bruto, str):
        texto = bruto
    else:
        texto = None
        for encoding in ('utf-8-sig', 'utf-8', 'cp1252', 'latin-1'):
            try:
                texto = bruto.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if texto is None:
            texto = bruto.decode('utf-8', errors='replace')
    # Extratos Bradesco usam só \r como quebra de linha.
    return texto.replace('\r\n', '\n').replace('\r', '\n')


def _iter_linhas_csv(texto):
    amostra = texto[:4096]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=';,')
    except csv.Error:
        dialeto = csv.excel
        dialeto.delimiter = ';' if amostra.count(';') >= amostra.count(',') else ','

    stream = io.StringIO(texto)
    leitor = csv.reader(stream, dialect=dialeto)
    cabecalho = None

    for indice, bruto in enumerate(leitor, start=1):
        valores = [(c or '').strip() for c in bruto]
        if not any(valores):
            continue

        if _parece_cabecalho(valores):
            cabecalho = [_norm_chave(v) for v in valores]
            continue

        if cabecalho:
            linha = {
                chave: valores[i] if i < len(valores) else ''
                for i, chave in enumerate(cabecalho)
                if chave
            }
        else:
            linha = _linha_lista(valores)

        if _linha_irrelevante(linha, valores):
            continue
        yield indice, linha


def _iter_linhas_planilha(arquivo):
    bruto = arquivo.read()
    if hasattr(bruto, 'encode'):
        bruto = bruto.encode('utf-8')
    try:
        wb = load_workbook(io.BytesIO(bruto), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(
            'Não foi possível ler a planilha. Use o XLSX exportado pela Porto Seguro.'
        ) from exc

    indice = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            indice += 1
            valores = [_celula_texto(c) for c in row]
            if not any(valores):
                continue
            if _parece_cabecalho_porto(valores):
                continue
            if _linha_porto_irrelevante(valores):
                continue
            linha = _linha_porto(valores)
            if linha:
                yield indice, linha


def _celula_texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _parece_cabecalho_porto(valores):
    texto = ' '.join(valores).lower()
    return 'crédito' in texto or 'credito' in texto


def _linha_porto_irrelevante(valores):
    if not valores:
        return True
    primeiro = (valores[0] or '').strip().upper()
    segundo = (valores[1] if len(valores) > 1 else '').strip().upper()
    if primeiro == 'TOTAL' or segundo == 'TOTAL':
        return True
    if segundo in {'PAGAMENTO', 'PAGAMENTO RECEBIDO'}:
        return True
    if not _parece_data(primeiro):
        return True
    return False


def _linha_porto(valores):
    """Porto: Data | Descrição | Crédito | Débito | Total | Cartão."""
    descricao = valores[1] if len(valores) > 1 else ''
    debito = valores[3] if len(valores) > 3 else ''
    if not descricao or not debito or debito in {'', ' '}:
        return None
    return {
        'descricao': descricao,
        'valor': debito,
        'origem': 'porto',
    }


def _parece_cabecalho(valores):
    texto = ' '.join(valores).lower()
    return (
        'historico' in texto
        or 'histórico' in texto
        or 'histуrico' in texto
        or ('descricao' in texto or 'descrição' in texto)
        or ('valor(r$)' in texto and 'data' in texto)
        or (texto.startswith('data') and 'valor' in texto)
        or ('date' in texto and 'title' in texto and 'amount' in texto)
        or ('date' in texto and 'amount' in texto)
    )


def _norm_chave(texto):
    return (texto or '').strip().lower()


def _linha_lista(valores):
    dados = {
        'descricao': valores[0] if len(valores) > 0 else '',
        'valor': valores[1] if len(valores) > 1 else '',
    }
    if len(valores) > 2:
        dados['quantidade_parcelas'] = valores[2]
    if len(valores) > 3:
        dados['parcela_atual'] = valores[3]
    # Formato posicional Bradesco sem cabeçalho: Data;Histórico;US$;R$
    if len(valores) >= 4 and _parece_data(valores[0]):
        dados = {
            'descricao': valores[1],
            'valor': valores[3],
        }
    return dados


def _parece_data(texto):
    return bool(re.match(r'^\d{1,2}/\d{1,2}(/\d{2,4})?$', (texto or '').strip()))


def _linha_irrelevante(linha, valores):
    juntos = ' '.join(valores).lower()
    if 'total da fatura' in juntos or juntos.startswith('resumo'):
        return True
    if juntos.startswith('taxas') or 'pagamento de contas' in juntos:
        return True
    if juntos.startswith('data:') or ('situa' in juntos and 'fatura' in juntos):
        return True
    if ';;;' in ''.join(valores):
        return True
    descricao = (
        _pegar(linha, COLUNAS_DESCRICAO)
        or linha.get('descricao', '')
        or (valores[1] if len(valores) > 1 else valores[0] if valores else '')
    )
    if IGNORAR_DESCRICAO.search((descricao or '').strip()):
        return True
    return False


def _linha_para_item(linha):
    if not linha:
        return None

    descricao = _pegar(linha, COLUNAS_DESCRICAO) or linha.get('descricao', '')
    valor_bruto = _pegar(linha, COLUNAS_VALOR) or linha.get('valor', '')
    qtd_bruto = _pegar(linha, COLUNAS_QTD) or linha.get('quantidade_parcelas', '')
    atual_bruto = _pegar(linha, COLUNAS_ATUAL) or linha.get('parcela_atual', '')
    origem = linha.get('origem', '')

    if not descricao or not valor_bruto:
        return None
    if IGNORAR_DESCRICAO.search(descricao.strip()):
        return None

    try:
        valor = _parse_decimal(valor_bruto)
    except ValueError:
        return None
    if valor <= 0:
        return None

    descricao, quantidade, atual = _extrair_parcelas(
        descricao, qtd_bruto, atual_bruto, origem=origem,
    )
    if quantidade < 1:
        raise ValueError('quantidade de parcelas inválida.')
    if atual < 1 or atual > quantidade:
        raise ValueError(f'parcela atual ({atual}) fora do total ({quantidade}).')

    return {
        'descricao': descricao[:200],
        'valor': valor,
        'quantidade_parcelas': quantidade,
        'parcela_atual': atual,
    }


def _extrair_parcelas(descricao, qtd_bruto, atual_bruto, origem=''):
    quantidade = _parse_int(qtd_bruto, padrao=None)
    atual = _parse_int(atual_bruto, padrao=None)
    limpa = re.sub(r'\s+', ' ', descricao.strip())

    match = PARCELA_NUBANK.match(limpa)
    if match:
        limpa = match.group('desc').strip()
        if atual is None:
            atual = int(match.group('atual'))
        if quantidade is None:
            quantidade = int(match.group('total'))
        return limpa, quantidade or 1, atual or 1

    match = PARCELA_NO_TEXTO.match(limpa)
    if match and _parcela_valida(match.group('atual'), match.group('total')):
        limpa = match.group('desc').strip()
        if atual is None:
            atual = int(match.group('atual'))
        if quantidade is None:
            quantidade = int(match.group('total'))
        return limpa, quantidade or 1, atual or 1

    if origem == 'porto' or (quantidade is None and atual is None):
        limpa, qtd_meio, atual_meio = _extrair_parcela_no_meio(limpa)
        if quantidade is None:
            quantidade = qtd_meio
        if atual is None:
            atual = atual_meio

    if origem == 'porto':
        limpa = _limpar_local_porto(limpa)

    return limpa, quantidade or 1, atual or 1


def _limpar_local_porto(descricao):
    """Remove o sufixo de localidade no padrão Porto: '... NATAL BR' / '... SAO PAULO BR'."""
    limpa = re.sub(r'\s+', ' ', descricao.strip())
    if not re.search(r'\s+BR\s*$', limpa, flags=re.IGNORECASE):
        return limpa
    sem_br = re.sub(r'\s+BR\s*$', '', limpa, flags=re.IGNORECASE).rstrip()
    partes = sem_br.rsplit(' ', 1)
    if len(partes) != 2:
        return sem_br
    resto, ultima = partes
    if not _parece_cidade(ultima):
        return sem_br
    partes2 = resto.rsplit(' ', 1)
    if len(partes2) == 2 and partes2[1].upper() in CIDADES_COMPOSTAS_INICIO:
        return partes2[0].rstrip()
    return resto.rstrip()


def _parece_cidade(token):
    if not token or not token.isupper():
        return False
    # Exige ao menos uma letra (evita códigos numéricos).
    return any(c.isalpha() for c in token)


def _extrair_parcela_no_meio(descricao):
    candidatos = []
    for match in PARCELA_NO_MEIO.finditer(descricao):
        if _parcela_valida(match.group('atual'), match.group('total')):
            candidatos.append(match)
    if not candidatos:
        return descricao, None, None
    # Prefere a última ocorrência (padrão Porto: estabelecimento + parcela + cidade).
    match = candidatos[-1]
    limpa = (descricao[:match.start()] + descricao[match.end():]).strip()
    limpa = re.sub(r'\s+', ' ', limpa)
    return limpa, int(match.group('total')), int(match.group('atual'))


def _parcela_valida(atual, total):
    try:
        atual_i = int(atual)
        total_i = int(total)
    except (TypeError, ValueError):
        return False
    return 1 <= atual_i <= total_i <= 48 and total_i >= 2


def _pegar(linha, aliases):
    for chave, valor in linha.items():
        if chave in aliases:
            return valor
    return ''


def _parse_decimal(texto):
    limpo = str(texto).strip()
    limpo = re.sub(r'[R$\s]', '', limpo, flags=re.IGNORECASE)
    if not limpo or limpo in {'.', '-', ','}:
        raise ValueError('valor inválido.')
    if ',' in limpo and '.' in limpo:
        if limpo.rfind(',') > limpo.rfind('.'):
            limpo = limpo.replace('.', '').replace(',', '.')
        else:
            limpo = limpo.replace(',', '')
    elif ',' in limpo:
        limpo = limpo.replace('.', '').replace(',', '.')
    try:
        return Decimal(limpo)
    except InvalidOperation as exc:
        raise ValueError(f'valor "{texto}" inválido.') from exc


def _parse_int(texto, padrao=1):
    if texto is None or str(texto).strip() == '':
        return padrao
    try:
        return int(str(texto).strip().split('/')[0])
    except (TypeError, ValueError) as exc:
        raise ValueError(f'número "{texto}" inválido.') from exc
